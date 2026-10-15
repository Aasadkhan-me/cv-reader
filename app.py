"""Resume ATS Score Analyzer - Streamlit + Gemini Flash."""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

MODEL_NAME = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
MAX_RESUME_CHARS = 20000

PROMPT_TEMPLATE = """You are an expert ATS (Applicant Tracking System) analyst and career coach.
Analyze the resume below{jd_clause}. Be honest, specific and strict.

Return ONLY valid JSON (no markdown fences) with exactly this structure:
{{
  "ats_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "section_scores": {{
    "formatting": <0-100>,
    "keywords": <0-100>,
    "experience": <0-100>,
    "education": <0-100>,
    "skills": <0-100>
  }},
  "strengths": ["<strength>", "..."],
  "improvements": [
    {{"area": "<section or topic>", "issue": "<what is wrong>", "suggestion": "<how to fix it>"}}
  ],
  "missing_keywords": ["<keyword>", "..."]
}}

Scoring guide: ATS-friendly structure and clear headings, relevant keywords,
quantified achievements, action verbs, contact details, consistent formatting,
and appropriate length.{jd_scoring}

{jd_block}RESUME:
\"\"\"
{resume}
\"\"\"
"""


# ----------------------------- Helpers ---------------------------------- #
def extract_text(uploaded_file) -> str:
    """Extract plain text from a PDF, DOCX or TXT upload."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


def build_prompt(resume_text: str, job_description: str = "") -> str:
    jd = job_description.strip()
    return PROMPT_TEMPLATE.format(
        jd_clause=" against the job description provided" if jd else "",
        jd_scoring=(
            " Weigh keyword match against the job description heavily."
            if jd
            else ""
        ),
        jd_block=f'JOB DESCRIPTION:\n"""\n{jd[:8000]}\n"""\n\n' if jd else "",
        resume=resume_text[:MAX_RESUME_CHARS],
    )


def parse_json_response(raw: str) -> dict:
    """Parse model output into a dict, tolerating code fences / extra text."""
    text = (raw or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, flags=re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise ValueError("The AI response was not valid JSON. Please try again.")


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def normalize_result(data: dict) -> dict:
    """Make sure every field exists and has a sane type/range."""
    scores = data.get("section_scores") or {}
    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "area": str(item.get("area", "General")),
                    "issue": str(item.get("issue", "")),
                    "suggestion": str(item.get("suggestion", "")),
                }
            )
        else:
            improvements.append({"area": "General", "issue": "", "suggestion": str(item)})
    return {
        "ats_score": _clamp(data.get("ats_score")),
        "summary": str(data.get("summary", "")),
        "section_scores": {
            k: _clamp(scores.get(k))
            for k in ("formatting", "keywords", "experience", "education", "skills")
        },
        "strengths": [str(s) for s in (data.get("strengths") or [])],
        "improvements": improvements,
        "missing_keywords": [str(k) for k in (data.get("missing_keywords") or [])],
    }


def analyze_resume(resume_text: str, job_description: str, api_key: str) -> dict:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return normalize_result(parse_json_response(response.text))


def get_api_key() -> str:
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:  # no secrets file locally
        key = ""
    return key or os.getenv("GEMINI_API_KEY", "")


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 60:
        return "Good - needs some work"
    return "Needs significant improvement"


# ------------------------------- UI ------------------------------------- #
def main():
    st.set_page_config(page_title="Resume ATS Analyzer", page_icon="📄", layout="centered")
    st.title("📄 Resume ATS Score Analyzer")
    st.caption(f"Upload your resume and get an ATS score with improvement tips. Powered by {MODEL_NAME}.")

    api_key = get_api_key()
    with st.sidebar:
        st.header("Settings")
        if api_key:
            st.success("API key loaded")
        else:
            api_key = st.text_input("Gemini API key", type="password",
                                    help="Get a free key at https://aistudio.google.com/apikey")
        st.markdown("Your resume is sent to Google's Gemini API for analysis and is not stored by this app.")

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional)",
        height=150,
        placeholder="Paste a job description to check how well your resume matches it...",
    )

    if st.button("Analyze Resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar.")
            st.stop()
        try:
            with st.spinner("Reading resume..."):
                text = extract_text(uploaded)
            if len(text) < 50:
                st.error("Couldn't read enough text. If your PDF is a scanned image, "
                         "please upload a text-based PDF or a DOCX file.")
                st.stop()
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(text, job_description, api_key)
        except Exception as exc:
            st.error(f"Something went wrong: {exc}")
            st.stop()

        score = result["ats_score"]
        st.divider()
        st.subheader("ATS Score")
        col1, col2 = st.columns([1, 2])
        col1.metric("Score", f"{score}/100")
        col2.progress(score / 100, text=score_label(score))
        if result["summary"]:
            st.write(result["summary"])

        st.subheader("Section Scores")
        cols = st.columns(5)
        for col, (name, val) in zip(cols, result["section_scores"].items()):
            col.metric(name.capitalize(), val)

        if result["strengths"]:
            st.subheader("✅ Strengths")
            for s in result["strengths"]:
                st.markdown(f"- {s}")

        if result["improvements"]:
            st.subheader("🛠️ Improvements")
            for i, imp in enumerate(result["improvements"], 1):
                with st.expander(f"{i}. {imp['area']}", expanded=i <= 3):
                    if imp["issue"]:
                        st.markdown(f"**Issue:** {imp['issue']}")
                    if imp["suggestion"]:
                        st.markdown(f"**Fix:** {imp['suggestion']}")

        if result["missing_keywords"]:
            st.subheader("🔑 Suggested Keywords")
            st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))


if __name__ == "__main__":
    main()
